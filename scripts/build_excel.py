# -*- coding: utf-8 -*-
"""
W5-A · 生成 Excel 多 Sheet 汇总表
输入：outputs/metrics.json + quality.json
输出：超市运营分析_汇总表.xlsx
"""
import sys, io, os, json, datetime
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import _paths as P

ROOT = P.ROOT
OUT = P.OUT
XLSX = os.path.join(ROOT, f"{P.prefix()}_汇总表.xlsx")

M = json.load(open(os.path.join(OUT, "metrics.json"), encoding="utf-8"))
Q = json.load(open(os.path.join(OUT, "quality.json"), encoding="utf-8"))
B = M["baseline"]

# ---------- 样式 ----------
H_FILL = PatternFill("solid", fgColor="1F3A5F")
H_FONT = Font(color="FFFFFF", bold=True, size=11)
T_FONT = Font(bold=True, size=13, color="1F3A5F")
SEC = Font(bold=True, size=11, color="FFFFFF")
SEC_FILL = PatternFill("solid", fgColor="2B6CB0")
NOTE = Font(size=10, color="5A6B7B", italic=True)
RED = Font(color="C53030", bold=True)
GRN = Font(color="2F855A", bold=True)
THIN = Side(style="thin", color="DDE4EA")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FMT_INT = '#,##0'
FMT_PCT = '0.00"%"'

wb = openpyxl.Workbook()

# 固定文档时间戳：让同一份数据每次生成的 .xlsx 字节一致，
# 避免重新构建后 git 出现无意义的二进制 diff
FIXED_TS = datetime.datetime(2016, 1, 1, 0, 0, 0)
wb.properties.created = FIXED_TS
wb.properties.modified = FIXED_TS


def head(ws, row, cols, widths=None):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.fill = H_FILL; cell.font = H_FONT; cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    return row + 1


def rows_out(ws, start, data, fmts=None, color_neg_col=None):
    r = start
    for d in data:
        for i, v in enumerate(d, 1):
            cell = ws.cell(row=r, column=i, value=v)
            cell.border = BORDER
            if fmts and i in fmts:
                cell.number_format = fmts[i]
                if isinstance(v, (int, float)):
                    cell.alignment = Alignment(horizontal="right")
            if color_neg_col and i == color_neg_col and isinstance(v, (int, float)):
                cell.font = RED if v < 0 else GRN
        r += 1
    return r


def title(ws, text, sub=None):
    ws["A1"] = text; ws["A1"].font = T_FONT
    if sub:
        ws["A2"] = sub; ws["A2"].font = NOTE
    return 4


# ============================================================
# Sheet 1 · 说明与口径
# ============================================================
ws = wb.active; ws.title = "说明与口径"
ws.column_dimensions["A"].width = 22; ws.column_dimensions["B"].width = 96
r = title(ws, "超市运营数据分析 · 汇总表", "数据区间 2016-01 ~ 2019-12 ｜ 生成方式：etl.py → analysis.py → build_excel.py（全流程可复现）")
items = [
    ("数据来源", "《2016—2019 年某超市运营数据》，含订单（9,959 行）、退货（1,163 行）、销售人员（6 行）三张表"),
    ("分析范围", f"{B['起止']} ｜ {B['订单数']:,} 个订单 ｜ {B['客户数']:,} 位客户 ｜ {B['产品数']:,} 个产品"),
    ("利润率口径", "利润率 = 利润 ÷ 销售额（汇总口径，非逐行平均）"),
    ("退货率口径", "退货率 = 退货订单数 ÷ 唯一订单数。退货表必须先去重：1,163 行 → 296 行；不去重会把退货率误算成 41.99%"),
    ("负利润处理", "2,041 行负利润为业务真实值，不是异常值，全部保留。若按统计规则剔除，将抹掉 983,895.91 元的真实亏损问题"),
    ("折扣档位", "数据仅含 7 档折扣：0 / 10% / 20% / 25% / 40% / 60% / 80%"),
    ("隐私处理", "客户编号已匿名化为 C0001–C0790"),
    ("已知局限", "① 订单 ID 前缀存在 CN/US 混用但国家字段全为「中国」，未用于地理划分；② 情景测算为静态上界，未考虑需求流失弹性"),
    ("", ""),
    ("Sheet 索引", "关键指标总览 / M1-经营趋势 / M2-折扣与利润 / M3-品类结构 / M4-客户价值 / M5-区域业绩 / M6-退货履约 / 数据质量台账"),
]
for k, v in items:
    ws.cell(row=r, column=1, value=k).font = Font(bold=True, size=10.5)
    ws.cell(row=r, column=2, value=v).font = Font(size=10.5)
    ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 28 if len(v) > 60 else 16
    r += 1

# ============================================================
# Sheet 2 · 关键指标总览
# ============================================================
ws = wb.create_sheet("关键指标总览")
r = title(ws, "关键指标总览（全量 2016—2019）")
r = head(ws, r, ["指标", "数值", "说明"], [26, 20, 58])
kpi_rows = [
    ("销售额（元）", B["销售额"], "四年累计"),
    ("利润（元）", B["利润"], "四年累计"),
    ("整体利润率", B["利润率"], "利润 ÷ 销售额"),
    ("订单数", B["订单数"], "唯一订单"),
    ("客户数", B["客户数"], "去重客户"),
    ("产品数", B["产品数"], "去重产品"),
    ("订单明细行", B["明细行"], "含一单多行"),
    ("真实退货率", M["M6_return"]["真实退货率"], "退货表去重后按订单口径"),
    ("折扣≥40% 销售额占比", M["M2_high_discount"]["销售额占比"], "只占不到两成销售额"),
    ("折扣≥40% 亏损占全网亏损比", M["M2_high_discount"]["亏损占全网亏损比"], "亏损的绝对主力"),
    ("净亏损客户数", M["M4_customer_stats"]["亏损客户数"], f"占 {M['M4_customer_stats']['亏损客户占比']}% 的客户"),
    ("净亏损客户合计亏损（元）", M["M4_customer_stats"]["亏损客户合计亏损"], "需要优先停止高折扣的名单"),
]
r = rows_out(ws, r, kpi_rows, fmts={2: FMT_INT})
ws.cell(row=r + 1, column=1, value="缺口提示").font = Font(bold=True, size=10.5)
ws.cell(row=r + 1, column=3, value="2019 年销售额同比 +28.72%，利润仅 +9.51%，利润率下降 2.19 个百分点 —— 增收不增利。").font = NOTE

# ============================================================
# Sheet 3 · M1 经营趋势
# ============================================================
ws = wb.create_sheet("M1-经营趋势")
r = title(ws, "M1 经营趋势：这四年生意变好了吗")
r = head(ws, r, ["年份", "销售额（元）", "利润（元）", "利润率", "订单数", "客单价（元）", "销售额同比", "利润同比"],
         [10, 16, 16, 11, 11, 14, 13, 12])
r = rows_out(ws, r, [(x["年"], x["销售额"], x["利润"], x["利润率"], x["订单数"], x["客单价"],
                      x["销售额同比"], x["利润同比"]) for x in M["M1_yearly"]],
             fmts={2: FMT_INT, 3: FMT_INT, 4: FMT_PCT, 5: FMT_INT, 6: FMT_INT, 7: FMT_PCT, 8: FMT_PCT},
             color_neg_col=3)

r += 1
ws.cell(row=r, column=1, value="利润率下滑归因（排除法）").font = SEC; ws.cell(row=r, column=1).fill = SEC_FILL
for c in range(2, 9):
    ws.cell(row=r, column=c).fill = SEC_FILL
r += 1
r = head(ws, r, ["项目", "2016", "2017", "2018", "2019", "结论"], [22, 12, 12, 12, 12, 46]) if False else r
for i, c in enumerate(["项目", "2016", "2017", "2018", "2019"], 1):
    cell = ws.cell(row=r, column=i, value=c); cell.font = Font(bold=True, size=10.5); cell.border = BORDER
r += 1
A = M["M1_attribution"]
attr_rows = [
    ("加权平均折扣（%）", *[A["加权平均折扣"][str(y)] for y in (2016, 2017, 2018, 2019)]),
    ("无折扣业务利润率（%）", *[A["无折扣业务利润率"][str(y)] for y in (2016, 2017, 2018, 2019)]),
    ("家具销售额占比（%）", *[A["类别销售额占比"][str(y)]["家具"] for y in (2016, 2017, 2018, 2019)]),
    ("家具利润率（%）", *[A["类别利润率"][str(y)]["家具"] for y in (2016, 2017, 2018, 2019)]),
    ("桌子利润（元）", *[A["桌子利润"][str(y)] for y in (2016, 2017, 2018, 2019)]),
]
r = rows_out(ws, r, attr_rows, fmts={2: '0.00', 3: '0.00', 4: '0.00', 5: '0.00'})
ws.cell(row=r, column=1, value="排除项").font = Font(bold=True, size=10.5)
ws.cell(row=r, column=2, value=A["排除项"]).font = Font(size=10)
r += 1
ws.cell(row=r, column=1, value="真因").font = Font(bold=True, size=10.5)
ws.cell(row=r, column=2, value=A["真因"]).font = Font(size=10, bold=True, color="C53030")

r += 2
ws.cell(row=r, column=1, value="季度与月度销售额").font = SEC; ws.cell(row=r, column=1).fill = SEC_FILL
for c in range(2, 5):
    ws.cell(row=r, column=c).fill = SEC_FILL
r += 1
r = head(ws, r, ["季度", "销售额（元）", "利润（元）", "利润率"], [10, 16, 16, 11])
r = rows_out(ws, r, [(x["季"], x["销售额"], x["利润"], x["利润率"]) for x in M["M1_quarterly"]],
             fmts={2: FMT_INT, 3: FMT_INT, 4: FMT_PCT}, color_neg_col=3)
r += 1
r = head(ws, r, ["月份", "销售额（元）", "利润（元）"], [10, 16, 16])
rows_out(ws, r, [(x["月"], x["销售额"], x["利润"]) for x in M["M1_monthly"]],
         fmts={2: FMT_INT, 3: FMT_INT}, color_neg_col=3)

# ============================================================
# Sheet 4 · M2 折扣与利润
# ============================================================
ws = wb.create_sheet("M2-折扣与利润")
r = title(ws, "M2 折扣与利润：钱被什么吃掉了")
r = head(ws, r, ["折扣档", "明细行", "销售额（元）", "销售额占比", "利润（元）", "利润率", "平均单笔销售额（元）"],
         [10, 11, 17, 12, 17, 11, 20])
r = rows_out(ws, r, [(x["折扣"], x["明细行"], x["销售额"], x["销售额占比"], x["利润"], x["利润率"], x["平均单笔销售额"])
                     for x in M["M2_discount"]],
             fmts={2: FMT_INT, 3: FMT_INT, 4: FMT_PCT, 5: FMT_INT, 6: FMT_PCT, 7: FMT_INT},
             color_neg_col=5)
r += 1
h = M["M2_high_discount"]
ws.cell(row=r, column=1, value="核心结论").font = Font(bold=True, size=11, color="C53030")
ws.cell(row=r, column=3, value=f"{h['阈值']} 的订单：占销售额 {h['销售额占比']}%、占明细行 {h['明细行占比']}%，"
                               f"却制造了全网 {h['亏损占全网亏损比']}% 的亏损（{h['利润']:,.0f} 元）").font = Font(size=10.5, bold=True)
r += 2
ws.cell(row=r, column=1, value="情景测算（静态上界，未计需求流失）").font = SEC; ws.cell(row=r, column=1).fill = SEC_FILL
for c in range(2, 6):
    ws.cell(row=r, column=c).fill = SEC_FILL
r += 1
r = head(ws, r, ["情景", "保留销售额（元）", "销售额变化率", "保留利润（元）", "利润变化率", "保留利润率"],
         [22, 18, 14, 18, 14, 12])
for key in ("scenario_40", "scenario_20"):
    s = M["M2_whatif"][key]
    r = rows_out(ws, r, [(s["剔除阈值"], s["保留销售额"], s["销售额变化率"], s["保留利润"], s["利润变化率"], s["保留利润率"])],
                 fmts={2: FMT_INT, 3: FMT_PCT, 4: FMT_INT, 5: FMT_PCT, 6: FMT_PCT}, color_neg_col=4)
r += 1
ws.cell(row=r, column=1, value="类别 × 折扣档 利润矩阵（元）").font = SEC; ws.cell(row=r, column=1).fill = SEC_FILL
for c in range(2, 9):
    ws.cell(row=r, column=c).fill = SEC_FILL
r += 1
disc_cols = ["类别"] + [f"{int(round(d*100))}%" for d in [0, 0.1, 0.2, 0.25, 0.4, 0.6, 0.8]]
r = head(ws, r, disc_cols, [12] + [12] * 7)
for row in M["M2_cat_discount_profit"]:
    r = rows_out(ws, r, [tuple([row["类别"]] + [row[c] for c in disc_cols[1:]])], fmts={i: FMT_INT for i in range(2, 9)})

# ============================================================
# Sheet 5 · M3 品类结构
# ============================================================
ws = wb.create_sheet("M3-品类结构")
r = title(ws, "M3 品类结构：哪些品类在失血")
r = head(ws, r, ["类别", "销售额（元）", "销售额占比", "利润（元）", "利润率", "平均折扣", "明细行"],
         [12, 16, 12, 16, 11, 11, 11])
r = rows_out(ws, r, [(x["类别"], x["销售额"], x["销售额占比"], x["利润"], x["利润率"], x["平均折扣"], x["明细行"])
                     for x in M["M3_category"]],
             fmts={2: FMT_INT, 3: FMT_PCT, 4: FMT_INT, 5: FMT_PCT, 6: FMT_PCT, 7: FMT_INT}, color_neg_col=4)
r += 1
qs = M["M3_quadrant_median"]
ws.cell(row=r, column=1, value=f"象限划分：销售额中位数 {qs['销售额中位数']:,.0f} 元、利润率中位数 {qs['利润率中位数']}%").font = NOTE
r += 2
r = head(ws, r, ["子类别", "类别", "象限", "销售额（元）", "利润（元）", "利润率", "平均折扣", "明细行"],
         [12, 10, 16, 16, 16, 11, 11, 11])
r = rows_out(ws, r, [(x["子类别"], x["类别"], x["象限"], x["销售额"], x["利润"], x["利润率"], x["平均折扣"], x["明细行"])
                     for x in M["M3_subcategory"]],
             fmts={4: FMT_INT, 5: FMT_INT, 6: FMT_PCT, 7: FMT_PCT, 8: FMT_INT}, color_neg_col=5)

# ============================================================
# Sheet 6 · M4 客户价值
# ============================================================
ws = wb.create_sheet("M4-客户价值")
r = title(ws, "M4 客户价值：谁值得投，谁在倒贴")
st = M["M4_customer_stats"]
r = head(ws, r, ["指标", "数值"], [26, 22])
r = rows_out(ws, r, [
    ("客户总数", st["客户总数"]), ("净亏损客户数", st["亏损客户数"]),
    ("净亏损客户占比", st["亏损客户占比"]), ("净亏损客户合计亏损（元）", st["亏损客户合计亏损"]),
    ("客均订单数", st["客均订单数"]), ("客均销售额（元）", st["客均销售额"]), ("客均利润（元）", st["客均利润"]),
    ("订单数中位数", st["订单数中位数"]), ("最多下单次数", st["订单数最多"]),
    ("仅下单 1 次的客户占比", st["一次下单客户占比"]),
    ("销售额 P50（元）", st["销售额分位"]["P50"]), ("销售额 P90（元）", st["销售额分位"]["P90"]),
], fmts={2: FMT_INT})
r += 1
ws.cell(row=r, column=1, value="说明：客户下单次数集中在中位数 6 次（区间 1–17），RFM 的 R/F 维度区分度不足，"
                               "故本模块改为「净亏损客户识别」，直接定位需要收紧折扣的客户名单。").font = NOTE
r += 2
r = head(ws, r, ["客户细分", "客户数", "订单数", "客单价（元）", "销售额（元）", "利润（元）", "利润率", "平均折扣"],
         [12, 10, 10, 13, 16, 16, 11, 11])
r = rows_out(ws, r, [(x["细分"], x["客户数"], x["订单数"], x["客单价"], x["销售额"], x["利润"], x["利润率"], x["平均折扣"])
                     for x in M["M4_segment"]],
             fmts={2: FMT_INT, 3: FMT_INT, 4: FMT_INT, 5: FMT_INT, 6: FMT_INT, 7: FMT_PCT, 8: FMT_PCT},
             color_neg_col=6)
r += 2
r = head(ws, r, ["客户编号", "订单数", "销售额（元）", "利润（元）", "利润率", "平均折扣"],
         [12, 10, 16, 16, 11, 11])
rows_out(ws, r, [(x["客户编号"], x["订单数"], x["销售额"], x["利润"], x["利润率"], x["平均折扣"])
                 for x in M["M4_loss_customers"]],
         fmts={2: FMT_INT, 3: FMT_INT, 4: FMT_INT, 5: FMT_PCT, 6: FMT_PCT}, color_neg_col=4)

# ============================================================
# Sheet 7 · M5 区域业绩
# ============================================================
ws = wb.create_sheet("M5-区域业绩")
r = title(ws, "M5 区域业绩：哪个大区该换打法")
r = head(ws, r, ["地区", "地区经理", "销售额（元）", "销售额占比", "利润（元）", "利润率", "订单数", "客户数", "平均折扣", "退货率"],
         [10, 11, 16, 12, 16, 11, 11, 11, 11, 11])
r = rows_out(ws, r, [(x["地区"], x["地区经理"], x["销售额"], x["销售额占比"], x["利润"], x["利润率"],
                      x["订单数"], x["客户数"], x["平均折扣"], x["退货率"]) for x in M["M5_region"]],
             fmts={3: FMT_INT, 4: FMT_PCT, 5: FMT_INT, 6: FMT_PCT, 7: FMT_INT, 8: FMT_INT, 9: FMT_PCT, 10: FMT_PCT},
             color_neg_col=5)
r += 2
ws.cell(row=r, column=1, value="各年度销售额（元）").font = SEC; ws.cell(row=r, column=1).fill = SEC_FILL
for c in range(2, 7):
    ws.cell(row=r, column=c).fill = SEC_FILL
r += 1
r = head(ws, r, ["地区", "2016", "2017", "2018", "2019"], [10, 15, 15, 15, 15])
r = rows_out(ws, r, [(x["地区"], x["2016"], x["2017"], x["2018"], x["2019"]) for x in M["M5_region_year"]],
             fmts={i: FMT_INT for i in range(2, 6)})
r += 1
r = head(ws, r, ["地区", "高折扣(≥40%)销售额占比"], [10, 22])
rows_out(ws, r, [(x["地区"], x["高折扣销售额占比"]) for x in M["M5_region_high_discount"]], fmts={2: FMT_PCT})

# ============================================================
# Sheet 8 · M6 退货履约
# ============================================================
ws = wb.create_sheet("M6-退货履约")
r = title(ws, "M6 退货与履约：退货是真问题还是噪音")
ret = M["M6_return"]
r = head(ws, r, ["项目", "数值"], [26, 22])
r = rows_out(ws, r, [
    ("退货表原始行数", ret["退货表原始行数"]),
    ("退货表去重后行数", ret["退货表去重后"]),
    ("删除的重复行数", ret["删除重复行"]),
    ("真实退货率（按订单）", ret["真实退货率"]),
    ("错误口径（按退货表行数）", ret["错误口径_按退货表行数"]),
], fmts={2: FMT_INT})
r += 1
ws.cell(row=r, column=1, value="注：不去重会得出 41.99% 的退货率，与真实的 10.69% 相差近 4 倍 —— 这是本项目最关键的一个口径陷阱。").font = NOTE
r += 2
r = head(ws, r, ["类别", "订单数", "退货率"], [12, 11, 11])
r = rows_out(ws, r, [(x["类别"], x["订单数"], x["退货率"]) for x in M["M6_return_by_category"]],
             fmts={2: FMT_INT, 3: FMT_PCT})
r += 1
r = head(ws, r, ["地区", "订单数", "退货率"], [12, 11, 11])
r = rows_out(ws, r, [(x["地区"], x["订单数"], x["退货率"]) for x in M["M6_return_by_region"]],
             fmts={2: FMT_INT, 3: FMT_PCT})
r += 1
r = head(ws, r, ["邮寄方式", "明细行", "平均发货天数", "销售额（元）", "利润率", "退货率"], [12, 11, 14, 16, 11, 11])
rows_out(ws, r, [(x["邮寄方式"], x["明细行"], x["平均发货天数"], x["销售额"], x["利润率"], x["退货率"])
                 for x in M["M6_shipping"]],
         fmts={2: FMT_INT, 3: '0.00', 4: FMT_INT, 5: FMT_PCT, 6: FMT_PCT})

# ============================================================
# Sheet 9 · 数据质量台账
# ============================================================
ws = wb.create_sheet("数据质量台账")
r = title(ws, "数据质量台账", f"数据源：{Q['数据源']}")
r = head(ws, r, ["检查项", "结果"], [34, 30])
q = Q["质量台账"]
order = ["退货表_原始行数", "退货表_去重后行数", "退货表_删除重复行数", "退货涉及唯一订单数",
         "订单表_明细行数", "订单表_唯一订单数", "订单表_唯一客户数", "订单表_唯一产品数",
         "订单表_全字段重复行", "订单表_缺失值总数", "订单表_负销售额行", "订单表_负数量行",
         "订单表_负利润行", "订单表_负利润金额", "订单表_发货早于下单行", "地区经理_缺失行",
         "真实退货率_按订单", "错误口径_按退货表行数"]
r = rows_out(ws, r, [(k, q[k]) for k in order], fmts={2: FMT_INT})
for rr in range(r - len(order), r):
    if ws.cell(row=rr, column=1).value in ("订单表_负利润行", "退货表_删除重复行数"):
        ws.cell(row=rr, column=2).font = RED
r += 1
ws.cell(row=r, column=1, value="国家/地区取值").font = Font(bold=True, size=10.5)
ws.cell(row=r, column=2, value="、".join(q["国家地区_取值"]))
r += 1
ws.cell(row=r, column=1, value="订单 ID 前缀分布").font = Font(bold=True, size=10.5)
ws.cell(row=r, column=2, value="、".join(f"{k} {v}" for k, v in q["订单ID_前缀分布"].items()))
r += 1
ws.cell(row=r, column=1, value="口径疑点说明").font = Font(bold=True, size=10.5)
ws.cell(row=r, column=2, value=q["订单ID_前缀异常说明"]).font = Font(size=10)
ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[r].height = 40

wb.save(XLSX)


def _normalize_zip(path):
    """让 .xlsx 输出字节完全确定。

    两处不确定来源：
    1. openpyxl 写出的 .xlsx 是 zip，条目时间戳取自当前时间；
    2. openpyxl 在 save 阶段会把 docProps/core.xml 里的 dcterms:modified
       重写为当前时间（即使事先设置了 wb.properties.modified）。
    不处理的话每次重跑都会产出内容相同、字节不同的文件，在 git 里表现为无意义的二进制 diff。
    """
    import zipfile, shutil, re
    fixed = b"2016-01-01T00:00:00Z"
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                              rb"\g<1>" + fixed + rb"\g<2>", data)
            item.date_time = (1980, 1, 1, 0, 0, 0)
            item.external_attr = 0o600 << 16
            zout.writestr(item, data)
    shutil.move(tmp, path)


_normalize_zip(XLSX)

print("=" * 72)
print("W5-A · Excel 汇总表生成完成")
print("=" * 72)
print(f"  输出：{XLSX}")
print(f"  Sheet：{len(wb.sheetnames)} 个 → {' / '.join(wb.sheetnames)}")
print(f"  体积：{os.path.getsize(XLSX)/1024:.1f} KB")
